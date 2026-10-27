import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

def create_presentation():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    NAVY = RGBColor(15, 23, 42)        # #0F172A
    CARD_BG = RGBColor(30, 41, 59)     # #1E293B
    LIGHT_BG = RGBColor(248, 250, 252) # #F8FAFC
    LIGHT_CARD = RGBColor(255, 255, 255)
    LIGHT_BORDER = RGBColor(226, 232, 240) # #E2E8F0
    TEXT_DARK = RGBColor(15, 23, 42)
    TEXT_MUTED = RGBColor(100, 116, 139) # #64748B
    TEXT_LIGHT = RGBColor(241, 245, 249)
    TEXT_LIGHT_MUTED = RGBColor(148, 163, 184)
    TEAL = RGBColor(14, 165, 233)       # #0EA5E9
    GREEN = RGBColor(34, 197, 94)       # #22C55E

    blank_layout = prs.slide_layouts[6]

    def add_header(slide, title_text, category_text, dark=False):
        tb_cat = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11), Inches(0.32))
        tf_cat = tb_cat.text_frame
        tf_cat.word_wrap = True
        tf_cat.margin_left = tf_cat.margin_right = tf_cat.margin_top = tf_cat.margin_bottom = 0
        p_cat = tf_cat.paragraphs[0]
        p_cat.text = category_text.upper()
        p_cat.font.name = "Segoe UI"
        p_cat.font.size = Pt(11)
        p_cat.font.bold = True
        p_cat.font.color.rgb = TEAL

        tb_title = slide.shapes.add_textbox(Inches(0.8), Inches(0.72), Inches(11.5), Inches(0.65))
        tf_title = tb_title.text_frame
        tf_title.word_wrap = True
        tf_title.margin_left = tf_title.margin_right = tf_title.margin_top = tf_title.margin_bottom = 0
        p_title = tf_title.paragraphs[0]
        p_title.text = title_text
        p_title.font.name = "Segoe UI"
        p_title.font.size = Pt(28)
        p_title.font.bold = True
        p_title.font.color.rgb = TEXT_LIGHT if dark else TEXT_DARK

    def add_takeaway(slide, text, slide_num, dark=False):
        line = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE,
            Inches(0.8), Inches(6.8), Inches(11.733), Inches(0.02)
        )
        line.fill.solid()
        line.fill.fore_color.rgb = RGBColor(51, 65, 85) if dark else RGBColor(226, 232, 240)
        line.line.color.rgb = line.fill.fore_color.rgb

        tb = slide.shapes.add_textbox(Inches(0.8), Inches(6.88), Inches(10.5), Inches(0.4))
        tf = tb.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        p = tf.paragraphs[0]
        run1 = p.add_run()
        run1.text = "KEY TAKEAWAY: "
        run1.font.bold = True
        run1.font.size = Pt(13)
        run1.font.name = "Segoe UI"
        run1.font.color.rgb = TEAL

        run2 = p.add_run()
        run2.text = text
        run2.font.bold = False
        run2.font.size = Pt(13)
        run2.font.name = "Segoe UI"
        run2.font.color.rgb = TEXT_LIGHT_MUTED if dark else TEXT_MUTED

        tb_num = slide.shapes.add_textbox(Inches(11.5), Inches(6.88), Inches(1.0), Inches(0.4))
        tf_num = tb_num.text_frame
        tf_num.margin_left = tf_num.margin_right = tf_num.margin_top = tf_num.margin_bottom = 0
        p_num = tf_num.paragraphs[0]
        p_num.alignment = PP_ALIGN.RIGHT
        p_num.text = f"{slide_num} / 7"
        p_num.font.bold = True
        p_num.font.size = Pt(13)
        p_num.font.name = "Segoe UI"
        p_num.font.color.rgb = TEXT_LIGHT_MUTED if dark else TEXT_MUTED

    # =========================================================================
    # SLIDE 1: Title & Problem (Dark Navy Theme)
    # =========================================================================
    slide1 = prs.slides.add_slide(blank_layout)
    bg1 = slide1.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
    bg1.fill.solid()
    bg1.fill.fore_color.rgb = NAVY
    bg1.line.fill.background()

    tb1_badge = slide1.shapes.add_textbox(Inches(0.8), Inches(0.8), Inches(5.5), Inches(0.4))
    tf1_badge = tb1_badge.text_frame
    tf1_badge.margin_left = tf1_badge.margin_top = 0
    p1_b = tf1_badge.paragraphs[0]
    p1_b.text = "ACADEMIC DBMS MICRO-PROJECT | ORACLE 23ai"
    p1_b.font.name = "Segoe UI"
    p1_b.font.size = Pt(11)
    p1_b.font.bold = True
    p1_b.font.color.rgb = TEAL

    tb1_title = slide1.shapes.add_textbox(Inches(0.8), Inches(1.3), Inches(5.5), Inches(2.2))
    tf1_title = tb1_title.text_frame
    tf1_title.word_wrap = True
    tf1_title.margin_left = tf1_title.margin_top = 0
    p1_t = tf1_title.paragraphs[0]
    p1_t.text = "ISP Management System"
    p1_t.font.name = "Segoe UI"
    p1_t.font.size = Pt(36)
    p1_t.font.bold = True
    p1_t.font.color.rgb = TEXT_LIGHT

    p1_sub = tf1_title.add_paragraph()
    p1_sub.text = 'Oracle Edition ("ISP Orbit")'
    p1_sub.font.name = "Segoe UI"
    p1_sub.font.size = Pt(22)
    p1_sub.font.bold = True
    p1_sub.font.color.rgb = TEAL

    p1_desc = tf1_title.add_paragraph()
    p1_desc.space_before = Pt(14)
    p1_desc.text = "A full-lifecycle telecom enterprise engine built with Oracle 23ai Free, PL/SQL packages, Pure JDBC, and JavaFX 21."
    p1_desc.font.name = "Segoe UI"
    p1_desc.font.size = Pt(15)
    p1_desc.font.color.rgb = TEXT_LIGHT_MUTED

    meta_box = slide1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(4.3), Inches(5.3), Inches(1.8))
    meta_box.fill.solid()
    meta_box.fill.fore_color.rgb = CARD_BG
    meta_box.line.color.rgb = RGBColor(51, 65, 85)
    tf1_meta = meta_box.text_frame
    tf1_meta.word_wrap = True
    tf1_meta.margin_left = tf1_meta.margin_top = Inches(0.25)
    p1_m1 = tf1_meta.paragraphs[0]
    p1_m1.text = "Presented by: Student Engineering Team"
    p1_m1.font.name = "Segoe UI"
    p1_m1.font.size = Pt(14)
    p1_m1.font.bold = True
    p1_m1.font.color.rgb = TEXT_LIGHT
    
    p1_m2 = tf1_meta.add_paragraph()
    p1_m2.text = "Course: Database Management Systems (CS-301)"
    p1_m2.font.name = "Segoe UI"
    p1_m2.font.size = Pt(13)
    p1_m2.font.color.rgb = TEXT_LIGHT_MUTED

    p1_m3 = tf1_meta.add_paragraph()
    p1_m3.text = "Target DB: Oracle Database 23ai Free (PDB: FREEPDB1)"
    p1_m3.font.name = "Segoe UI"
    p1_m3.font.size = Pt(13)
    p1_m3.font.color.rgb = TEAL

    tb1_prob_title = slide1.shapes.add_textbox(Inches(6.6), Inches(0.8), Inches(6.0), Inches(0.4))
    tf1_pt = tb1_prob_title.text_frame
    tf1_pt.margin_left = tf1_pt.margin_top = 0
    p1_pt = tf1_pt.paragraphs[0]
    p1_pt.text = "THE 4 REAL-WORLD TELECOM CHALLENGES WE SOLVE"
    p1_pt.font.name = "Segoe UI"
    p1_pt.font.size = Pt(12)
    p1_pt.font.bold = True
    p1_pt.font.color.rgb = TEAL

    problems = [
        ("Multi-Tier Recurring Billing", "Automates complex monthly cycles, data cap tiers, and 18% GST tax calculation without human error."),
        ("Concurrent Payment Race Conditions", "Eliminates double-credits and dirty reads when simultaneous payments hit the same invoice."),
        ("Fragmented Telecom Lifecycle", "Bridges customer onboarding, router hardware assignment, technician dispatch, and support tickets."),
        ("Compliance & Audit Trails", "Tracks immutable historical changes with pre/post JSON snapshots for financial non-repudiation.")
    ]

    for i, (p_title, p_body) in enumerate(problems):
        y_pos = Inches(1.35 + i * 1.3)
        p_card = slide1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.6), y_pos, Inches(5.9), Inches(1.15))
        p_card.fill.solid()
        p_card.fill.fore_color.rgb = CARD_BG
        p_card.line.color.rgb = RGBColor(51, 65, 85)
        tf_pc = p_card.text_frame
        tf_pc.word_wrap = True
        tf_pc.margin_left = tf_pc.margin_top = Inches(0.2)
        tf_pc.margin_right = Inches(0.2)
        p_item1 = tf_pc.paragraphs[0]
        p_item1.text = f"{i+1}. {p_title}"
        p_item1.font.name = "Segoe UI"
        p_item1.font.size = Pt(14)
        p_item1.font.bold = True
        p_item1.font.color.rgb = TEXT_LIGHT

        p_item2 = tf_pc.add_paragraph()
        p_item2.text = p_body
        p_item2.font.name = "Segoe UI"
        p_item2.font.size = Pt(12)
        p_item2.font.color.rgb = TEXT_LIGHT_MUTED

    add_takeaway(slide1, "Connects business operational workflows directly to database transaction guarantees.", 1, dark=True)
    slide1.notes_slide.notes_text_frame.text = (
        "Good morning everyone. Today we are presenting our DBMS micro-project: the ISP Management System, Oracle Edition. "
        "Modern Internet Service Providers face four critical database challenges: automated multi-tier billing, payment race conditions, "
        "managing the entire subscriber-to-technician lifecycle, and non-negotiable financial compliance. "
        "In this presentation, we demonstrate how Oracle 23ai and enterprise database patterns solve each of these problems reliably."
    )

    # =========================================================================
    # SLIDE 2: Architecture & Tech Stack (Light Theme)
    # =========================================================================
    slide2 = prs.slides.add_slide(blank_layout)
    bg2 = slide2.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
    bg2.fill.solid()
    bg2.fill.fore_color.rgb = LIGHT_BG
    bg2.line.fill.background()

    add_header(slide2, "End-to-End System Architecture", "System Design & Rationale")

    lbl_diag = slide2.shapes.add_textbox(Inches(0.8), Inches(1.5), Inches(5.2), Inches(0.35))
    tf_ld = lbl_diag.text_frame
    tf_ld.margin_left = tf_ld.margin_top = 0
    p_ld = tf_ld.paragraphs[0]
    p_ld.text = "3-LAYER SYSTEM INTEGRATION"
    p_ld.font.name = "Segoe UI"
    p_ld.font.size = Pt(12)
    p_ld.font.bold = True
    p_ld.font.color.rgb = TEXT_MUTED

    layers = [
        ("Desktop Presentation Layer", "JavaFX 21 (Controls + Graphics)", "Multithreaded Task<T> UI • Zero Freezes • Real-time Search", RGBColor(238, 242, 255), RGBColor(99, 102, 241)),
        ("Database Connectivity Layer", "Pure JDBC Driver (ojdbc11 23ai)", "Explicit ACID Transactions • Direct CallableStatement Execution", RGBColor(240, 253, 250), RGBColor(20, 184, 166)),
        ("Relational Storage & Logic", "Oracle Database 23ai (FREEPDB1)", "17 Tables (3NF) • PL/SQL Packages • Compound Triggers", RGBColor(254, 242, 242), RGBColor(239, 68, 68))
    ]

    for i, (l_role, l_tech, l_sub, l_bg, l_border) in enumerate(layers):
        y = Inches(1.9 + i * 1.5)
        box = slide2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), y, Inches(5.2), Inches(1.15))
        box.fill.solid()
        box.fill.fore_color.rgb = l_bg
        box.line.color.rgb = l_border
        box.line.width = Pt(1.5)
        tf_b = box.text_frame
        tf_b.word_wrap = True
        tf_b.margin_left = tf_b.margin_top = Inches(0.15)
        
        p1 = tf_b.paragraphs[0]
        p1.text = l_role.upper()
        p1.font.name = "Segoe UI"
        p1.font.size = Pt(10)
        p1.font.bold = True
        p1.font.color.rgb = TEXT_MUTED

        p2 = tf_b.add_paragraph()
        p2.text = l_tech
        p2.font.name = "Segoe UI"
        p2.font.size = Pt(14)
        p2.font.bold = True
        p2.font.color.rgb = TEXT_DARK

        p3 = tf_b.add_paragraph()
        p3.text = l_sub
        p3.font.name = "Segoe UI"
        p3.font.size = Pt(11)
        p3.font.color.rgb = TEXT_MUTED

        if i < 2:
            arrow = slide2.shapes.add_shape(MSO_SHAPE.DOWN_ARROW, Inches(3.2), y + Inches(1.18), Inches(0.35), Inches(0.28))
            arrow.fill.solid()
            arrow.fill.fore_color.rgb = TEAL
            arrow.line.fill.background()

    lbl_tab = slide2.shapes.add_textbox(Inches(6.4), Inches(1.5), Inches(6.1), Inches(0.35))
    tf_lt = lbl_tab.text_frame
    tf_lt.margin_left = tf_lt.margin_top = 0
    p_lt = tf_lt.paragraphs[0]
    p_lt.text = "CORE COMPONENT SELECTION & JUSTIFICATION"
    p_lt.font.name = "Segoe UI"
    p_lt.font.size = Pt(12)
    p_lt.font.bold = True
    p_lt.font.color.rgb = TEXT_MUTED

    tech_table_shape = slide2.shapes.add_table(5, 3, Inches(6.4), Inches(1.9), Inches(6.1), Inches(4.5))
    tech_table = tech_table_shape.table
    tech_table.columns[0].width = Inches(1.5)
    tech_table.columns[1].width = Inches(1.3)
    tech_table.columns[2].width = Inches(3.3)

    table_data = [
        ("Layer / Component", "Technology", "Design Justification (The 'Why')"),
        ("Java Platform", "Java 17 / 25 LTS", "High-throughput concurrency primitives and long-term platform stability."),
        ("Build Tool", "Apache Maven 3.9+", "Automated dependency lifecycle, artifact packaging, and build reproduction."),
        ("Credential Security", "jBCrypt 0.4", "Adaptive 12-round salted hashing preventing rainbow table lookups."),
        ("System Logging", "SLF4J + Logback", "Structured diagnostic tracing of JDBC statements and background thread tasks.")
    ]

    for row_idx, row in enumerate(table_data):
        for col_idx, cell_value in enumerate(row):
            cell = tech_table.cell(row_idx, col_idx)
            cell.text = cell_value
            p = cell.text_frame.paragraphs[0]
            p.font.name = "Segoe UI"
            if row_idx == 0:
                p.font.bold = True
                p.font.size = Pt(12)
                p.font.color.rgb = TEXT_LIGHT
                cell.fill.solid()
                cell.fill.fore_color.rgb = NAVY
            else:
                p.font.size = Pt(11)
                p.font.color.rgb = TEXT_DARK
                cell.fill.solid()
                cell.fill.fore_color.rgb = LIGHT_CARD if row_idx % 2 == 1 else RGBColor(241, 245, 249)

    add_takeaway(slide2, "Pure JDBC ensures zero ORM overhead and direct control over Oracle transactions.", 2, dark=False)
    slide2.notes_slide.notes_text_frame.text = (
        "Here is our system architecture. It follows a clean 3-layer enterprise pattern. "
        "At the top is a JavaFX 21 desktop application running non-blocking background workers so the interface never freezes. "
        "In the middle, we intentionally chose Pure JDBC over heavy ORM frameworks to give us explicit control over SQL transactions and stored procedure calls. "
        "At the foundation is Oracle Database 23ai Free, encapsulating data integrity in the database engine."
    )

    # =========================================================================
    # SLIDE 3: Database Design (Light Theme)
    # =========================================================================
    slide3 = prs.slides.add_slide(blank_layout)
    bg3 = slide3.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
    bg3.fill.solid()
    bg3.fill.fore_color.rgb = LIGHT_BG
    bg3.line.fill.background()

    add_header(slide3, "Relational Schema & 17 Core Tables", "Relational Database Modeling")

    clusters = [
        ("Customers & Locations", ["CUSTOMER", "ADDRESS", "SERVICE_AREA"], RGBColor(238, 242, 255), RGBColor(99, 102, 241)),
        ("Plans & Subscriptions", ["PLAN", "SUBSCRIPTION", "EQUIPMENT", "USAGE_RECORD"], RGBColor(240, 253, 244), RGBColor(34, 197, 94)),
        ("Billing & Revenue", ["INVOICE", "INVOICE_ITEM", "PAYMENT"], RGBColor(255, 247, 237), RGBColor(249, 115, 22)),
        ("Operations & Security", ["TECHNICIAN", "INSTALLATION", "SUPPORT_TICKET", "TICKET_UPDATE", "NETWORK_OUTAGE", "APP_USER", "AUDIT_LOG"], RGBColor(254, 242, 242), RGBColor(239, 68, 68))
    ]

    for i, (c_name, c_tables, c_bg, c_border) in enumerate(clusters):
        x = Inches(0.8 + i * 2.97)
        c_box = slide3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, Inches(1.5), Inches(2.85), Inches(2.0))
        c_box.fill.solid()
        c_box.fill.fore_color.rgb = c_bg
        c_box.line.color.rgb = c_border
        c_box.line.width = Pt(1.5)
        tf_c = c_box.text_frame
        tf_c.word_wrap = True
        tf_c.margin_left = tf_c.margin_top = Inches(0.15)
        
        p1 = tf_c.paragraphs[0]
        p1.text = c_name
        p1.font.name = "Segoe UI"
        p1.font.size = Pt(12)
        p1.font.bold = True
        p1.font.color.rgb = c_border

        p2 = tf_c.add_paragraph()
        p2.space_before = Pt(6)
        p2.text = f"{len(c_tables)} Tables:\n" + ", ".join(c_tables)
        p2.font.name = "Segoe UI"
        p2.font.size = Pt(11)
        p2.font.color.rgb = TEXT_DARK

    er_card = slide3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(3.7), Inches(3.75), Inches(2.85))
    er_card.fill.solid()
    er_card.fill.fore_color.rgb = LIGHT_CARD
    er_card.line.color.rgb = LIGHT_BORDER
    tf_er = er_card.text_frame
    tf_er.word_wrap = True
    tf_er.margin_left = tf_er.margin_top = Inches(0.2)
    p_er1 = tf_er.paragraphs[0]
    p_er1.text = "CORE ER RELATIONSHIPS"
    p_er1.font.bold = True
    p_er1.font.size = Pt(12)
    p_er1.font.color.rgb = TEAL

    er_bullets = [
        "CUSTOMER 1:N SUBSCRIPTION (Contracts)",
        "PLAN 1:N SUBSCRIPTION (Catalog links)",
        "SUBSCRIPTION 1:N INVOICE (Monthly billing)",
        "INVOICE 1:N PAYMENT (Payment ledger)",
        "TECHNICIAN 1:N TICKET (Support log)"
    ]
    for b in er_bullets:
        p = tf_er.add_paragraph()
        p.space_before = Pt(4)
        p.text = f"• {b}"
        p.font.size = Pt(11)
        p.font.color.rgb = TEXT_DARK

    norm_card = slide3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(4.78), Inches(3.7), Inches(3.75), Inches(2.85))
    norm_card.fill.solid()
    norm_card.fill.fore_color.rgb = LIGHT_CARD
    norm_card.line.color.rgb = LIGHT_BORDER
    tf_n = norm_card.text_frame
    tf_n.word_wrap = True
    tf_n.margin_left = tf_n.margin_top = Inches(0.2)
    p_n1 = tf_n.paragraphs[0]
    p_n1.text = "NORMALIZATION & INDEXING"
    p_n1.font.bold = True
    p_n1.font.size = Pt(12)
    p_n1.font.color.rgb = TEAL

    norm_bullets = [
        "3NF/BCNF Model: Zero transitive dependencies or insertion/update anomalies.",
        "Foreign Key Indexes: idx_cust_area, idx_sub_cust eliminate table scans.",
        "Composite Indexes: idx_inv_status_due on (status, due_date) speeds up queries.",
        "Search Indexes: Case-insensitive lookups on customer names and phone numbers."
    ]
    for b in norm_bullets:
        p = tf_n.add_paragraph()
        p.space_before = Pt(4)
        p.text = f"• {b}"
        p.font.size = Pt(11)
        p.font.color.rgb = TEXT_DARK

    views_card = slide3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(8.76), Inches(3.7), Inches(3.75), Inches(2.85))
    views_card.fill.solid()
    views_card.fill.fore_color.rgb = LIGHT_CARD
    views_card.line.color.rgb = LIGHT_BORDER
    tf_v = views_card.text_frame
    tf_v.word_wrap = True
    tf_v.margin_left = tf_v.margin_top = Inches(0.2)
    p_v1 = tf_v.paragraphs[0]
    p_v1.text = "6 ANALYTICAL VIEWS"
    p_v1.font.bold = True
    p_v1.font.size = Pt(12)
    p_v1.font.color.rgb = TEAL

    views_bullets = [
        "V_ACTIVE_SUBSCRIPTIONS: Live customer plan registry",
        "V_OVERDUE_INVOICES: Aging receivables & risk flag",
        "V_CUSTOMER_BALANCE: Lifetime billed vs. paid balance",
        "V_MONTHLY_REVENUE: Monthly fiscal plan breakdown",
        "V_TECHNICIAN_WORKLOAD: Real-time ticket dispatch status",
        "V_DATA_USAGE_SUMMARY: Quota consumption & FUP"
    ]
    for b in views_bullets:
        p = tf_v.add_paragraph()
        p.space_before = Pt(4)
        p.text = f"• {b}"
        p.font.size = Pt(11)
        p.font.color.rgb = TEXT_DARK

    add_takeaway(slide3, "3NF structure eliminates data redundancy while composite indexes deliver instant analytical queries.", 3, dark=False)
    slide3.notes_slide.notes_text_frame.text = (
        "Slide 3 shows our database schema. For those new to DBMS terminology, 3NF means Third Normal Form—every piece of data is stored in exactly one place so there is no duplicate data. "
        "We organized all 17 tables into four clear business clusters: Customers, Plans, Billing, and Operations. "
        "Notice the composite index idx_inv_status_due: by indexing both status and due date together, the database instantly finds overdue accounts without scanning millions of records. "
        "We also created 6 dedicated views to power the dashboard and analytics tab without complex joins in UI code."
    )

    # =========================================================================
    # SLIDE 4: PL/SQL Engine (Light Theme - Independent Shapes, No Overlap)
    # =========================================================================
    slide4 = prs.slides.add_slide(blank_layout)
    bg4 = slide4.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
    bg4.fill.solid()
    bg4.fill.fore_color.rgb = LIGHT_BG
    bg4.line.fill.background()

    add_header(slide4, "PL/SQL Package: PKG_BILLING & Concurrency Control", "Database Engine Logic")

    # Left: PKG_BILLING Procedures Card
    left_card = slide4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.5), Inches(5.6), Inches(5.05))
    left_card.fill.solid()
    left_card.fill.fore_color.rgb = LIGHT_CARD
    left_card.line.color.rgb = LIGHT_BORDER
    tf_lc = left_card.text_frame
    tf_lc.word_wrap = True
    tf_lc.margin_left = tf_lc.margin_top = Inches(0.25)
    
    p_lp1 = tf_lc.paragraphs[0]
    p_lp1.text = "CORE PACKAGE PROCEDURES & FUNCTIONS"
    p_lp1.font.bold = True
    p_lp1.font.size = Pt(13)
    p_lp1.font.color.rgb = TEAL

    pkg_items = [
        ("sp_record_payment", "Applies pessimistic row lock (FOR UPDATE NOWAIT) on invoice, inserts payment ledger record, and auto-updates status."),
        ("sp_generate_monthly_invoices", "Batch generates monthly recurring invoices for active subscriptions with automatic 18% GST calculation."),
        ("sp_suspend_overdue_accounts", "Automatically flags 30+ day delinquent accounts and suspends active subscriptions to eliminate risk."),
        ("fn_calculate_late_fee", "Deterministic function applying a 2% compound overdue fee penalty per overdue billing cycle.")
    ]

    for p_name, p_desc in pkg_items:
        p_head = tf_lc.add_paragraph()
        p_head.space_before = Pt(8)
        p_head.text = f"• {p_name}()"
        p_head.font.bold = True
        p_head.font.size = Pt(12)
        p_head.font.color.rgb = TEXT_DARK

        p_body = tf_lc.add_paragraph()
        p_body.text = p_desc
        p_body.font.size = Pt(11)
        p_body.font.color.rgb = TEXT_MUTED

    # Right: Pipeline Card (Top Half)
    rt_top = slide4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.8), Inches(1.5), Inches(5.7), Inches(3.1))
    rt_top.fill.solid()
    rt_top.fill.fore_color.rgb = LIGHT_CARD
    rt_top.line.color.rgb = LIGHT_BORDER
    tf_rt = rt_top.text_frame
    tf_rt.word_wrap = True
    tf_rt.margin_left = tf_rt.margin_top = Inches(0.22)

    p_rp1 = tf_rt.paragraphs[0]
    p_rp1.text = "SAFE PAYMENT TRANSACTION PIPELINE"
    p_rp1.font.bold = True
    p_rp1.font.size = Pt(13)
    p_rp1.font.color.rgb = TEAL

    steps = [
        ("Step 1: Lock Invoice Row", "Pessimistic lock (FOR UPDATE NOWAIT) prevents double payments."),
        ("Step 2: Ledger Insertion", "Inserts payment audit entry with timestamp & transaction mode."),
        ("Step 3: Status Transition", "Calculates new balance -> status becomes PAID or PARTIALLY_PAID.")
    ]

    for s_title, s_sub in steps:
        p_st = tf_rt.add_paragraph()
        p_st.space_before = Pt(6)
        p_st.text = f"✔ {s_title}"
        p_st.font.bold = True
        p_st.font.size = Pt(11.5)
        p_st.font.color.rgb = GREEN

        p_sb = tf_rt.add_paragraph()
        p_sb.text = s_sub
        p_sb.font.size = Pt(10.5)
        p_sb.font.color.rgb = TEXT_MUTED

    # Right: Code Box (Bottom Half)
    code_box = slide4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.8), Inches(4.75), Inches(5.7), Inches(1.8))
    code_box.fill.solid()
    code_box.fill.fore_color.rgb = NAVY
    code_box.line.color.rgb = RGBColor(51, 65, 85)
    tf_code = code_box.text_frame
    tf_code.word_wrap = True
    tf_code.margin_left = tf_code.margin_top = Inches(0.2)

    sql_lines = [
        "-- Pessimistic Row Locking in sp_record_payment",
        "SELECT total_amount, amount_paid INTO v_total, v_paid",
        "  FROM invoice WHERE invoice_id = p_invoice_id",
        " FOR UPDATE NOWAIT;"
    ]
    for j, line in enumerate(sql_lines):
        p_line = tf_code.paragraphs[0] if j == 0 else tf_code.add_paragraph()
        p_line.text = line
        p_line.font.name = "Consolas"
        p_line.font.size = Pt(10)
        p_line.font.color.rgb = TEAL if j == 0 or "FOR UPDATE" in line else RGBColor(226, 232, 240)

    add_takeaway(slide4, "PL/SQL pessimistic row locking prevents payment race conditions and guarantees financial accuracy.", 4, dark=False)
    slide4.notes_slide.notes_text_frame.text = (
        "In Slide 4, we examine PKG_BILLING, our PL/SQL package. "
        "A critical concept here is pessimistic locking using FOR UPDATE NOWAIT. "
        "If two users or payment gateways attempt to pay the exact same invoice at the same second, Oracle locks the row for the first transaction, preventing duplicate credits or dirty reads. "
        "The package also automates monthly batch invoice generation with 18% GST and automatically suspends accounts overdue past 30 days."
    )

    # =========================================================================
    # SLIDE 5: Triggers and Audit (Light Theme - Independent Shapes, No Overlap)
    # =========================================================================
    slide5 = prs.slides.add_slide(blank_layout)
    bg5 = slide5.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
    bg5.fill.solid()
    bg5.fill.fore_color.rgb = LIGHT_BG
    bg5.line.fill.background()

    add_header(slide5, "Compound Triggers & Non-Repudiation Auditing", "Database Integrity & Event Automation")

    # Left: Mutating Table Card (Top Half)
    mut_top = slide5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.5), Inches(5.6), Inches(3.1))
    mut_top.fill.solid()
    mut_top.fill.fore_color.rgb = LIGHT_CARD
    mut_top.line.color.rgb = LIGHT_BORDER
    tf_mt = mut_top.text_frame
    tf_mt.word_wrap = True
    tf_mt.margin_left = tf_mt.margin_top = Inches(0.2)
    
    p_m1 = tf_mt.paragraphs[0]
    p_m1.text = "MUTATING TABLE PROBLEM & SOLUTION"
    p_m1.font.bold = True
    p_m1.font.size = Pt(13)
    p_m1.font.color.rgb = TEAL

    p_m2 = tf_mt.add_paragraph()
    p_m2.space_before = Pt(4)
    p_m2.text = "• The ORA-04091 Anomaly:"
    p_m2.font.bold = True
    p_m2.font.size = Pt(11)
    p_m2.font.color.rgb = RGBColor(239, 68, 68)

    p_m3 = tf_mt.add_paragraph()
    p_m3.text = "If a row trigger queries the table being modified, Oracle throws ORA-04091: table is mutating (like inspecting an engine while driving)."
    p_m3.font.size = Pt(10)
    p_m3.font.color.rgb = TEXT_DARK

    p_m4 = tf_mt.add_paragraph()
    p_m4.space_before = Pt(4)
    p_m4.text = "• The Compound Trigger Solution:"
    p_m4.font.bold = True
    p_m4.font.size = Pt(11)
    p_m4.font.color.rgb = GREEN

    p_m5 = tf_mt.add_paragraph()
    p_m5.text = "trg_invoice_status_compound separates execution into two safe phases:\n1. Row-Level: Collects modified IDs in memory array.\n2. AFTER STATEMENT: Safely evaluates statuses."
    p_m5.font.size = Pt(10)
    p_m5.font.color.rgb = TEXT_DARK

    # Left: Trigger Code Box (Bottom Half)
    c_box = slide5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(4.75), Inches(5.6), Inches(1.8))
    c_box.fill.solid()
    c_box.fill.fore_color.rgb = NAVY
    c_box.line.color.rgb = RGBColor(51, 65, 85)
    tf_cb = c_box.text_frame
    tf_cb.word_wrap = True
    tf_cb.margin_left = tf_cb.margin_top = Inches(0.18)
    code_trg = [
        "FOR EACH ROW -- Phase 1: In-memory collection",
        "  g_invoice_ids.EXTEND; g_invoice_ids(LAST) := :NEW.invoice_id;",
        "AFTER STATEMENT -- Phase 2: Safe execution",
        "  UPDATE invoice ... WHERE invoice_id IN (SELECT ...);"
    ]
    for k, line in enumerate(code_trg):
        p_c = tf_cb.paragraphs[0] if k == 0 else tf_cb.add_paragraph()
        p_c.text = line
        p_c.font.name = "Consolas"
        p_c.font.size = Pt(9.5)
        p_c.font.color.rgb = TEAL if "--" in line else RGBColor(226, 232, 240)

    # Right: Audit Card (Top Half)
    aud_top = slide5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.8), Inches(1.5), Inches(5.7), Inches(3.1))
    aud_top.fill.solid()
    aud_top.fill.fore_color.rgb = LIGHT_CARD
    aud_top.line.color.rgb = LIGHT_BORDER
    tf_at = aud_top.text_frame
    tf_at.word_wrap = True
    tf_at.margin_left = tf_at.margin_top = Inches(0.2)

    p_a1 = tf_at.paragraphs[0]
    p_a1.text = "IMMUTABLE AUDIT TRAIL (AUDIT_LOG)"
    p_a1.font.bold = True
    p_a1.font.size = Pt(13)
    p_a1.font.color.rgb = TEAL

    audit_features = [
        ("Automated Event Capture", "Triggers trg_customer_audit and trg_payment_audit automatically record all DML events."),
        ("Pre & Post JSON Snapshots", "Stores :OLD and :NEW row states as structured JSON strings for non-repudiation."),
        ("Regulatory Compliance", "Audit logs are strictly write-only, meeting enterprise compliance standards.")
    ]

    for a_t, a_d in audit_features:
        p_at = tf_at.add_paragraph()
        p_at.space_before = Pt(5)
        p_at.text = f"✔ {a_t}"
        p_at.font.bold = True
        p_at.font.size = Pt(11)
        p_at.font.color.rgb = TEXT_DARK

        p_ad = tf_at.add_paragraph()
        p_ad.text = a_d
        p_ad.font.size = Pt(10)
        p_ad.font.color.rgb = TEXT_MUTED

    # Right: Audit JSON Box (Bottom Half)
    j_box = slide5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.8), Inches(4.75), Inches(5.7), Inches(1.8))
    j_box.fill.solid()
    j_box.fill.fore_color.rgb = NAVY
    j_box.line.color.rgb = RGBColor(51, 65, 85)
    tf_jb = j_box.text_frame
    tf_jb.word_wrap = True
    tf_jb.margin_left = tf_jb.margin_top = Inches(0.18)
    json_lines = [
        "-- Sample AUDIT_LOG Change Capture",
        'OLD_STATE: {"status": "ACTIVE", "plan_id": 101}',
        'NEW_STATE: {"status": "SUSPENDED", "plan_id": 101}',
        'ACTION: "UPDATE" | USER: "isp_app" | TIMESTAMP'
    ]
    for k, line in enumerate(json_lines):
        p_j = tf_jb.paragraphs[0] if k == 0 else tf_jb.add_paragraph()
        p_j.text = line
        p_j.font.name = "Consolas"
        p_j.font.size = Pt(9.5)
        p_j.font.color.rgb = TEAL if k == 0 else RGBColor(226, 232, 240)

    add_takeaway(slide5, "Compound triggers eliminate mutating-table errors; JSON audit trails ensure full compliance.", 5, dark=False)
    slide5.notes_slide.notes_text_frame.text = (
        "Slide 5 covers an advanced Oracle concept: the mutating-table error, or ORA-04091. "
        "In plain English, if a trigger tries to read a table while that same table is still being updated, Oracle stops it—just like you cannot safely work on an engine while the car is moving. "
        "We solved this using a Compound Trigger: during the row phase, we only record changed IDs into memory; then in AFTER STATEMENT, we safely update the table without conflict. "
        "We also implemented automated audit triggers that capture before and after data as JSON snapshots for financial non-repudiation."
    )

    # =========================================================================
    # SLIDE 6: Security and the JavaFX App (Light Theme - Independent Shapes, No Overlap)
    # =========================================================================
    slide6 = prs.slides.add_slide(blank_layout)
    bg6 = slide6.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
    bg6.fill.solid()
    bg6.fill.fore_color.rgb = LIGHT_BG
    bg6.line.fill.background()

    add_header(slide6, "Database Security & Multithreaded JavaFX Client", "Security & Application Design")

    # Left: Security Card (Top Half)
    sec_top = slide6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.5), Inches(5.6), Inches(2.9))
    sec_top.fill.solid()
    sec_top.fill.fore_color.rgb = LIGHT_CARD
    sec_top.line.color.rgb = LIGHT_BORDER
    tf_st = sec_top.text_frame
    tf_st.word_wrap = True
    tf_st.margin_left = tf_st.margin_top = Inches(0.2)

    p_s1 = tf_st.paragraphs[0]
    p_s1.text = "DEFENSE-IN-DEPTH SECURITY"
    p_s1.font.bold = True
    p_s1.font.size = Pt(13)
    p_s1.font.color.rgb = TEAL

    sec_items = [
        ("jBCrypt Password Hashing", "Adaptive cost factor 12 salted hashing protects credentials in APP_USER from rainbow tables."),
        ("Oracle Security Profile", "isp_app_profile locks accounts after 5 failed logins with 90-day password expiration."),
        ("Least Privilege Role", "isp_readonly_role restricts analytical reporting tools from executing write operations.")
    ]

    for st, sd in sec_items:
        p_st = tf_st.add_paragraph()
        p_st.space_before = Pt(4)
        p_st.text = f"✔ {st}"
        p_st.font.bold = True
        p_st.font.size = Pt(11)
        p_st.font.color.rgb = TEXT_DARK

        p_sd = tf_st.add_paragraph()
        p_sd.text = sd
        p_sd.font.size = Pt(10)
        p_sd.font.color.rgb = TEXT_MUTED

    # Left: RBAC Table (Bottom Half)
    rbac_table_shape = slide6.shapes.add_table(5, 2, Inches(0.8), Inches(4.55), Inches(5.6), Inches(2.0))
    rbac_table = rbac_table_shape.table
    rbac_table.columns[0].width = Inches(1.8)
    rbac_table.columns[1].width = Inches(3.8)

    rbac_data = [
        ("System Role (RBAC)", "Authorized Application Modules"),
        ("ADMIN", "Full Access: All 7 Modules & Configuration"),
        ("BILLING_CLERK", "Billing, Payments, Invoices, Customer Lookup"),
        ("TECHNICIAN", "Assigned Installations, Hardware, Tech Tickets"),
        ("SUPPORT_AGENT", "Customer Profiles, Ticket Entry, Work Logs")
    ]
    for r_idx, (r_role, r_mods) in enumerate(rbac_data):
        c0 = rbac_table.cell(r_idx, 0)
        c1 = rbac_table.cell(r_idx, 1)
        c0.text = r_role
        c1.text = r_mods
        for c in (c0, c1):
            p = c.text_frame.paragraphs[0]
            p.font.name = "Segoe UI"
            if r_idx == 0:
                p.font.bold = True
                p.font.size = Pt(10)
                p.font.color.rgb = TEXT_LIGHT
                c.fill.solid()
                c.fill.fore_color.rgb = NAVY
            else:
                p.font.size = Pt(9.5)
                p.font.color.rgb = TEXT_DARK
                c.fill.solid()
                c.fill.fore_color.rgb = LIGHT_CARD if r_idx % 2 == 1 else RGBColor(241, 245, 249)

    # Right: JavaFX Desktop Client Engineering Highlights Card (Full Height)
    app_card = slide6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.8), Inches(1.5), Inches(5.7), Inches(5.05))
    app_card.fill.solid()
    app_card.fill.fore_color.rgb = LIGHT_CARD
    app_card.line.color.rgb = LIGHT_BORDER
    tf_ap = app_card.text_frame
    tf_ap.word_wrap = True
    tf_ap.margin_left = tf_ap.margin_top = Inches(0.25)

    p_ap1 = tf_ap.paragraphs[0]
    p_ap1.text = "JAVAFX CLIENT ENGINEERING HIGHLIGHTS"
    p_ap1.font.bold = True
    p_ap1.font.size = Pt(13)
    p_ap1.font.color.rgb = TEAL

    app_highlights = [
        ("Non-Blocking UI Threads", "All SQL operations run in JavaFX Task<T> worker threads. The GUI remains 100% responsive without freeze."),
        ("Atomic Multi-Step Wizard", "The subscription provisioning wizard wraps 3 DML operations inside an explicit transaction (commit/rollback)."),
        ("Oracle Error Trapping", "Catches PL/SQL ORA-20001 through ORA-20005 exceptions and displays clean, operator-friendly dialogs."),
        ("Modern UX Architecture", "Real-time client search (Ctrl+F), unified modals (X, Cancel, ESC), Dark/Light mode, and 1-click CSV exports.")
    ]

    for h_t, h_d in app_highlights:
        p_ht = tf_ap.add_paragraph()
        p_ht.space_before = Pt(8)
        p_ht.text = f"• {h_t}"
        p_ht.font.bold = True
        p_ht.font.size = Pt(12)
        p_ht.font.color.rgb = TEXT_DARK

        p_hd = tf_ap.add_paragraph()
        p_hd.text = h_d
        p_hd.font.size = Pt(11)
        p_hd.font.color.rgb = TEXT_MUTED

    add_takeaway(slide6, "Role-based security protects enterprise data; non-blocking JavaFX guarantees zero UI freezes.", 6, dark=False)
    slide6.notes_slide.notes_text_frame.text = (
        "Slide 6 presents our security model and desktop client engineering. "
        "RBAC stands for Role-Based Access Control: an employee only sees the screens their role permits. For example, a technician cannot touch financial invoices, and a support agent cannot edit network plans. "
        "User passwords are protected with jBCrypt at cost factor 12. "
        "On the frontend, all database communication runs on asynchronous background threads, ensuring the desktop application never stutters or freezes during heavy database operations."
    )

    # =========================================================================
    # SLIDE 7: Testing, Results & Future Scope (Dark Navy Theme)
    # =========================================================================
    slide7 = prs.slides.add_slide(blank_layout)
    bg7 = slide7.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
    bg7.fill.solid()
    bg7.fill.fore_color.rgb = NAVY
    bg7.line.fill.background()

    add_header(slide7, "Validation, Results & Future Horizons", "Project Verification & Summary", dark=True)

    card_test = slide7.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.5), Inches(5.6), Inches(5.05))
    card_test.fill.solid()
    card_test.fill.fore_color.rgb = CARD_BG
    card_test.line.color.rgb = RGBColor(51, 65, 85)
    tf_t = card_test.text_frame
    tf_t.word_wrap = True
    tf_t.margin_left = tf_t.margin_top = Inches(0.25)

    p_t1 = tf_t.paragraphs[0]
    p_t1.text = "RIGOROUS SYSTEM VERIFICATION"
    p_t1.font.bold = True
    p_t1.font.size = Pt(13)
    p_t1.font.color.rgb = TEAL

    test_results = [
        ("Schema & DDL Integrity", "Zero compilation errors across 17 tables, 6 views, 8 triggers, and PKG_BILLING package."),
        ("Concurrency Lock Testing", "Verified FOR UPDATE NOWAIT row lock under concurrent payment posts with zero race conditions or deadlocks."),
        ("Soft-Delete Referencing", "Verified subscriber deactivations preserve referential integrity across historical invoices and payments."),
        ("High-Volume Seed Dataset", "Loaded 20+ complete seed customer profiles with multi-criteria real-time querying and CSV exports.")
    ]

    for tr_t, tr_d in test_results:
        p_trt = tf_t.add_paragraph()
        p_trt.space_before = Pt(8)
        p_trt.text = f"✔ {tr_t}"
        p_trt.font.bold = True
        p_trt.font.size = Pt(12)
        p_trt.font.color.rgb = GREEN

        p_trd = tf_t.add_paragraph()
        p_trd.text = tr_d
        p_trd.font.size = Pt(11)
        p_trd.font.color.rgb = TEXT_LIGHT_MUTED

    card_fut = slide7.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.8), Inches(1.5), Inches(5.7), Inches(5.05))
    card_fut.fill.solid()
    card_fut.fill.fore_color.rgb = CARD_BG
    card_fut.line.color.rgb = RGBColor(51, 65, 85)
    tf_f = card_fut.text_frame
    tf_f.word_wrap = True
    tf_f.margin_left = tf_f.margin_top = Inches(0.25)

    p_f1 = tf_f.paragraphs[0]
    p_f1.text = "FUTURE ENHANCEMENTS & ROADMAP"
    p_f1.font.bold = True
    p_f1.font.size = Pt(13)
    p_f1.font.color.rgb = TEAL

    future_scope = [
        ("1. Self-Service Customer Portal", "Web-based portal for subscribers to monitor real-time data caps, download PDF bills, and pay via UPI."),
        ("2. Automated SNMP Network Monitoring", "Automated network trap listeners logging grid faults directly to NETWORK_OUTAGE."),
        ("3. Payment Gateway Webhooks", "Live Razorpay and Stripe API integration calling sp_record_payment in real-time.")
    ]

    for fs_t, fs_d in future_scope:
        p_fst = tf_f.add_paragraph()
        p_fst.space_before = Pt(8)
        p_fst.text = fs_t
        p_fst.font.bold = True
        p_fst.font.size = Pt(12)
        p_fst.font.color.rgb = TEXT_LIGHT

        p_fsd = tf_f.add_paragraph()
        p_fsd.text = fs_d
        p_fsd.font.size = Pt(11)
        p_fsd.font.color.rgb = TEXT_LIGHT_MUTED

    qa_box = slide7.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(7.05), Inches(4.9), Inches(5.2), Inches(1.4))
    qa_box.fill.solid()
    qa_box.fill.fore_color.rgb = RGBColor(15, 23, 42)
    qa_box.line.color.rgb = TEAL
    qa_box.line.width = Pt(1.5)
    tf_qa = qa_box.text_frame
    tf_qa.word_wrap = True
    tf_qa.margin_left = tf_qa.margin_top = Inches(0.18)

    p_qa1 = tf_qa.paragraphs[0]
    p_qa1.alignment = PP_ALIGN.CENTER
    p_qa1.text = "THANK YOU!"
    p_qa1.font.bold = True
    p_qa1.font.size = Pt(18)
    p_qa1.font.color.rgb = TEAL

    p_qa2 = tf_qa.add_paragraph()
    p_qa2.alignment = PP_ALIGN.CENTER
    p_qa2.text = "Questions & Discussion are welcome."
    p_qa2.font.size = Pt(13)
    p_qa2.font.color.rgb = TEXT_LIGHT_MUTED

    add_takeaway(slide7, "Proves that robust enterprise database architecture ensures bulletproof software applications.", 7, dark=True)
    slide7.notes_slide.notes_text_frame.text = (
        "To conclude, our testing validated all 17 tables, PL/SQL packages, and compound triggers with zero errors. "
        "We proved that concurrent payments are safely handled without deadlocks, and soft-deletes preserve historical audit records. "
        "Looking forward, this database foundation can seamlessly support a customer-facing web portal, SNMP network monitoring, and online payment gateways. "
        "Thank you for your time, and we are now open for any questions."
    )

    output_path = os.path.join(os.getcwd(), "ISP_Management_System_Presentation.pptx")
    prs.save(output_path)
    print(f"Presentation successfully created at: {output_path}")

if __name__ == "__main__":
    create_presentation()
